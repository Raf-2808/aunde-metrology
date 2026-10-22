"""Odczyt danych XLSM/XLSX bez uruchamiania VBA i bez zaufania do formuł."""
import io,hashlib,zipfile
from sqlalchemy import select
import openpyxl
from .models import HistoricalRow

def importuj_historie(db,user,content):
    """Zachowuje niepełne i jakościowe dane jako historię do przeglądu."""
    z=zipfile.ZipFile(io.BytesIO(content))
    if sum(x.file_size for x in z.infolist())>100*1024*1024: raise ValueError('Skoroszyt po rozpakowaniu przekracza 100 MB')
    wb=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
    if 'statistic' not in wb.sheetnames: raise ValueError('Brak arkusza statistic')
    ws=wb['statistic'];digest=hashlib.sha256(content).hexdigest();inserted=0;skipped=0
    if ws.max_row>100000 or ws.max_column>200: raise ValueError('Przekroczony rozmiar tabeli')
    for i,row in enumerate(ws.iter_rows(min_row=4,values_only=True),4):
        if not row[0]:continue
        if db.scalar(select(HistoricalRow.id).where(HistoricalRow.file_hash==digest,HistoricalRow.row_number==i)):
            skipped+=1;continue
        values={openpyxl.utils.get_column_letter(j+1):v.isoformat() if hasattr(v,'isoformat') else v for j,v in enumerate(row) if v is not None}
        issues=['Niezweryfikowana historyczna wersja tolerancji','Rekord historyczny nie jest nową zatwierdzoną sesją']
        if any(isinstance(v,str) and v.startswith('=') for v in row):issues.append('Formuły zachowano jako tekst; nie wykonywano')
        if any(v in ('NR','-') or v is None for v in row[4:53]):issues.append('Brakujące lub pominięte wyniki')
        db.add(HistoricalRow(file_hash=digest,row_number=i,lot=str(row[0]),values=values,issues=issues,author=user.id));inserted+=1
    wb.close();return {'inserted':inserted,'skipped':skipped,'sha256':digest,'status':'HISTORICAL_REVIEW'}
