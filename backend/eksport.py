"""Lokalne generowanie XLSX i obrazów QR bez usług zewnętrznych."""

from io import BytesIO
from datetime import date
import qrcode
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Alignment

KOLUMNY = [
    ("Lp", None),
    ("Numer", "numer"),
    ("Typ", "typ"),
    ("Stan", "stan"),
    ("Status", "status"),
    ("Wydział", "wydzial"),
    ("Grupa", "klasa"),
    ("R&R", "rr"),
    ("Ostatni", "ostatni"),
    ("Interwał", "interwal"),
    ("Następny", "nastepny"),
    ("Stan techniczny", "stan_techniczny"),
    ("Status metrologiczny", "metrologiczny"),
    ("Dopuszczenie", "dopuszczenie"),
    ("Jednostka interwału", "jednostka"),
    ("Komentarz do dat", "komentarz_dat"),
    ("Ostatni odczyt QR", "ostatni_qr"),
    ("Liczba odczytów QR", "liczba_qr"),
    ("Adres karty QR", None),
    ("Kod QR", None),
]


# Generuje XLSX z osobną kolumną adresu i ostatnią kolumną obrazu QR; nie zmienia tokenów ani ewidencji.
def zbuduj_excel(rekordy: list[dict], tokeny: dict[str, str], adres: str) -> bytes:
    ksiazka = Workbook()
    arkusz = ksiazka.active
    arkusz.title = "Przyrządy"
    arkusz.append([n for n, _ in KOLUMNY])
    obrazy = []
    for lp, rekord in enumerate(rekordy, 1):
        wartosci = []
        for nazwa, pole in KOLUMNY:
            wartosc = rekord.get(pole) if pole else lp if nazwa == "Lp" else ""
            if pole in ("ostatni", "nastepny") and wartosc:
                wartosc = date.fromisoformat(wartosc)
            wartosci.append(wartosc)
        arkusz.append(wartosci)
        for kolumna, wartosc in enumerate(wartosci, 1):
            komorka = arkusz.cell(lp + 1, kolumna)
            if isinstance(wartosc, str):
                komorka.data_type = "s"
            if isinstance(wartosc, date):
                komorka.number_format = "yyyy-mm-dd"
            komorka.alignment = Alignment(vertical="center", wrap_text=True)
        token = tokeny.get(rekord["id"])
        if not token:
            raise ValueError("Brak aktualnego tokenu QR dla przyrządu. Eksport przerwany")
        url = adres.rstrip("/") + "/p/" + token
        qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=6, border=4)
        qr.add_data(url)
        qr.make(fit=True)
        bufor = BytesIO()
        qr.make_image(fill_color="black", back_color="white").save(bufor, format="PNG")
        bufor.seek(0)
        obrazy.append(bufor)
        obraz = Image(bufor)
        obraz.width = obraz.height = 160
        arkusz.add_image(obraz, f"T{lp + 1}")
        komorka = arkusz.cell(lp + 1, 19, url)
        komorka.hyperlink = url
        komorka.alignment = Alignment(vertical="center", wrap_text=True)
        arkusz.row_dimensions[lp + 1].height = 155
    for komorka in arkusz[1]:
        komorka.fill = PatternFill("solid", fgColor="17324D")
        komorka.font = Font(color="FFFFFF", bold=True)
        komorka.alignment = Alignment(vertical="center", wrap_text=True)
    from openpyxl.utils import get_column_letter

    for nr in range(1, len(KOLUMNY) + 1):
        arkusz.column_dimensions[get_column_letter(nr)].width = (
            50 if nr == 19 else 25 if nr in (2, 3, 13, 20) else 18
        )
    arkusz.freeze_panes = "C2"
    arkusz.auto_filter.ref = arkusz.dimensions
    arkusz.sheet_view.showGridLines = False
    arkusz.print_options.horizontalCentered = True
    arkusz.print_title_rows = "1:1"
    wynik = BytesIO()
    ksiazka.save(wynik)
    return wynik.getvalue()
