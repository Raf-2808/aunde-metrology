"""Polecenia IT: konta, uprawnienia bazodanowe, import próbny i kontrola terminów."""

import argparse
from getpass import getpass
import hashlib
import json
from pathlib import Path
import os
import smtplib
import ssl
from email.message import EmailMessage
from datetime import date, datetime, timedelta
from argon2 import PasswordHasher
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError
from .baza import SesjaBazy, silnik
from .modele import (
    Uzytkownik,
    Import,
    WierszImportu,
    Harmonogram,
    Przyrzad,
    DaneRR,
    Powiadomienie,
    Audyt,
    teraz_utc,
)
from .ewidencja import dodaj_przyrzad, pobierz_rejestr, dzis_warszawa
from .reguly import sprawdz_daty

NAGLOWKI = [
    "Lp",
    "Numer",
    "Typ",
    "Stan",
    "Status",
    "Wydział",
    "Grupa",
    "R&R",
    "Ostatni",
    "Interwał",
    "Następny",
]


# Nadaje osobnym rolom aplikacji minimalne uprawnienia; wymaga połączenia roli migracyjnej.
def ustaw_uprawnienia() -> None:
    with silnik.begin() as polaczenie:
        for rola, zmienna in (
            ("metrologia_admin", "ADMIN_DB_PASSWORD"),
            ("metrologia_public", "PUBLIC_DB_PASSWORD"),
        ):
            haslo = os.environ[zmienna]
            if len(haslo) < 24 or not haslo.isalnum() or not haslo.isascii():
                raise ValueError("Hasła ról bazy muszą mieć co najmniej 24 znaki alfanumeryczne")
            # Nazwy ról są stałe, hasło ograniczone do znaków alfanumerycznych przed DDL.
            try:
                if not polaczenie.scalar(
                    text("SELECT 1 FROM pg_roles WHERE rolname=:rola"), {"rola": rola}
                ):
                    polaczenie.exec_driver_sql(f"CREATE ROLE {rola} LOGIN PASSWORD '{haslo}'")
                else:
                    polaczenie.exec_driver_sql(f"ALTER ROLE {rola} PASSWORD '{haslo}'")
                polaczenie.exec_driver_sql(f"ALTER ROLE {rola} SET timezone='UTC'")
            except DBAPIError:
                raise RuntimeError(
                    "Nie udało się skonfigurować roli bazy. Sprawdź uprawnienia roli migracyjnej"
                ) from None
        polaczenie.exec_driver_sql("REVOKE CREATE ON SCHEMA public FROM PUBLIC")
        polaczenie.exec_driver_sql(
            "GRANT USAGE ON SCHEMA public TO metrologia_admin,metrologia_public"
        )
        polaczenie.exec_driver_sql(
            "GRANT SELECT,INSERT,UPDATE ON ALL TABLES IN SCHEMA public TO metrologia_admin"
        )
        polaczenie.exec_driver_sql("GRANT DELETE ON public.sesje TO metrologia_admin")
        polaczenie.exec_driver_sql(
            "REVOKE UPDATE ON public.odczyty_qr,public.audyt,public.historia_statusow FROM metrologia_admin"
        )
        polaczenie.exec_driver_sql(
            "REVOKE ALL ON ALL TABLES IN SCHEMA public FROM metrologia_public"
        )
        polaczenie.exec_driver_sql(
            "GRANT EXECUTE ON FUNCTION public.przygotuj_karte(text),public.zapisz_odczyt(text,text,jsonb),public.historia_karty(text,integer) TO metrologia_public"
        )


# Tworzy pierwszego Administratora z hasłem wprowadzanym interaktywnie; nie przyjmuje hasła w argv.
def utworz_administratora(login: str) -> None:
    haslo = getpass("Hasło (min. 12 znaków): ")
    if len(haslo) < 12 or haslo != getpass("Powtórz hasło: "):
        raise ValueError("Hasło jest za krótkie lub potwierdzenie jest inne")
    with SesjaBazy.begin() as baza:
        if baza.scalar(select(Uzytkownik.id).where(Uzytkownik.login == login)):
            raise ValueError("Konto już istnieje")
        baza.add(
            Uzytkownik(
                login=login,
                skrot_hasla=PasswordHasher().hash(haslo),
                rola="Administrator",
            )
        )


# Odczytuje rzeczywistą datę Excela; tekst daty zerowej zamienia na brak, bez zgadywania.
def odczytaj_date(wartosc) -> date | None:
    if isinstance(wartosc, datetime):
        return wartosc.date()
    if isinstance(wartosc, date):
        return wartosc
    if wartosc in (None, "", "0000-00-00"):
        return None
    return date.fromisoformat(str(wartosc))


# Analizuje XLSX lub importuje atomowo; scala wyłącznie numery z jawnej konfiguracji poza repozytorium.
def importuj_excel(
    sciezka: str, zapisz: bool, raport: str, polacz_numery: set[str] | None = None
) -> dict:
    from openpyxl import load_workbook

    plik = Path(sciezka)
    ksiazka = load_workbook(plik, data_only=False)
    if ksiazka.sheetnames != ["Przyrządy"] or [k.value for k in ksiazka.active[4]] != NAGLOWKI:
        raise ValueError(
            "Nieobsługiwana struktura arkusza. Wymagane zatwierdzenie nowego mapowania"
        )
    dane = []
    problemy = []
    numery: dict[str, list] = {}
    for nr, wiersz in enumerate(ksiazka.active.iter_rows(min_row=5, values_only=True), 5):
        r = dict(zip(NAGLOWKI, wiersz))
        bledy = []
        try:
            if any(isinstance(v, str) and v.startswith("=") for v in wiersz):
                raise ValueError("Formuła w danych źródłowych wymaga analizy")
            r["data"] = odczytaj_date(r["Ostatni"])
            sprawdz_daty(r["Status"], r["data"], dzis_warszawa())
            if r["Status"] not in ("Aktywna", "Nieaktywna") or r["Stan"] not in (
                "Na stanie",
                "Zwrócona",
            ):
                raise ValueError("Nieznany status lub stan")
            if r["R&R"] not in ("---", "Tak"):
                raise ValueError("Nieznana wartość R&R")
            from .reguly import oblicz_termin, normalizuj_wydzial

            r["wydzial"] = normalizuj_wydzial(str(r["Wydział"]))
            if not isinstance(r["Interwał"], int) or r["Interwał"] <= 0:
                raise ValueError("Nieprawidłowy interwał")
            nastepny = odczytaj_date(r["Następny"])
            if r["data"] and nastepny != oblicz_termin(r["data"], r["Interwał"], "miesiac"):
                raise ValueError("Termin źródłowy różni się od obliczanego")
        except (ValueError, TypeError, OverflowError) as blad:
            bledy.append(str(blad))
        if bledy:
            problemy.append({"wiersz": nr, "numer": r["Numer"], "bledy": bledy})
        r["wiersz"] = nr
        dane.append(r)
        numery.setdefault(str(r["Numer"]), []).append(r)
    duplikaty = {n: [r["wiersz"] for r in rr] for n, rr in numery.items() if len(rr) > 1}
    if set(duplikaty) - (polacz_numery or set()):
        problemy.append({"bledy": ["Nowe duplikaty numerów wymagają uzgodnienia"]})
    wynik = {
        "wiersze": len(dane),
        "przyrzady": len(numery),
        "duplikaty": duplikaty,
        "problemy": problemy,
        "zapisano": False,
    }
    Path(raport).write_text(json.dumps(wynik, ensure_ascii=False, indent=2), encoding="utf-8")
    if not zapisz:
        return wynik
    if problemy:
        raise ValueError(
            "Import wstrzymany. Popraw wyjątki w kopii źródła według raportu; oryginał pozostaw bez zmian"
        )
    skrot = hashlib.sha256(plik.read_bytes()).hexdigest()
    with SesjaBazy.begin() as baza:
        if baza.scalar(select(Import.id).where(Import.sha256 == skrot)):
            return {**wynik, "juz_zaimportowano": True}
        if baza.scalar(select(Przyrzad.id).limit(1)):
            raise ValueError(
                "Automatyczny import pierwszorazowy wymaga pustego rejestru. Kolejny plik wymaga uzgodnienia UUID"
            )
        partia = Import(sha256=skrot)
        baza.add(partia)
        baza.flush()
        for numer, wiersze in numery.items():
            r = max(wiersze, key=lambda x: (x["data"] or date.min, x["wiersz"]))
            p = dodaj_przyrzad(baza, numer, str(r["Typ"]), r["wydzial"], None)
            p.stan, p.status = r["Stan"], r["Status"]
            p.komentarz_dat = sprawdz_daty(p.status, r["data"], dzis_warszawa())
            h = baza.scalar(
                select(Harmonogram).where(
                    Harmonogram.przyrzad_id == p.id, Harmonogram.proces == "sprawdzenie"
                )
            )
            h.ostatni, h.interwal = r["data"], r["Interwał"]
            h.termin_zrodlowy = odczytaj_date(r["Następny"])
            baza.get(DaneRR, p.id).wartosc = r["R&R"]
            for oryginal in wiersze:
                surowe = {
                    k: v.isoformat() if isinstance(v, (date, datetime)) else v
                    for k, v in oryginal.items()
                    if k in NAGLOWKI
                }
                baza.add(
                    WierszImportu(
                        import_id=partia.id,
                        wiersz=oryginal["wiersz"],
                        przyrzad_id=p.id,
                        surowe=surowe,
                        problemy=[],
                    )
                )
            baza.add(
                Audyt(
                    rekord_id=p.id,
                    operacja="Import",
                    szczegoly={"partia": partia.id, "klasa_poczatkowa": "B"},
                )
            )
    wynik["zapisano"] = True
    Path(raport).write_text(json.dumps(wynik, ensure_ascii=False, indent=2), encoding="utf-8")
    return wynik


# Ocenia przekroczenie wymaganego terminu niezależnie od innych przyczyn blokady, bez skutków ubocznych.
def czy_wymaga_alertu(rekord: dict, dzis: date) -> bool:
    return (
        rekord.get("status") == "Aktywna"
        and rekord.get("klasa") in ("A", "B")
        and bool(rekord.get("nastepny"))
        and date.fromisoformat(rekord["nastepny"]) < dzis
    )


# Kolejkuje alerty terminu i wysyła SMTP po ponownej ocenie pod blokadą przyrządu; błędy planują ponowienie.
# Przy EMAIL_ENABLED=false nie wysyła ani nie kolejkuje. Zwraca None; blokada użytkowania działa niezależnie.
def obsluz_powiadomienia() -> None:
    if os.getenv("EMAIL_ENABLED", "false").lower() != "true":
        return
    with SesjaBazy.begin() as baza:
        if not baza.scalar(text("SELECT pg_try_advisory_xact_lock(13420600)")):
            return
        dane = pobierz_rejestr(baza)
        for r in dane:
            if not czy_wymaga_alertu(r, dzis_warszawa()):
                continue
            klucz = f"{r['id']}:{r['proces']}:{r['nastepny']}"
            oczekujace = baza.scalar(select(Powiadomienie).where(Powiadomienie.klucz == klucz))
            if oczekujace is None:
                baza.add(
                    Powiadomienie(
                        klucz=klucz,
                        przyrzad_id=r["id"],
                        dane={k: r[k] for k in ("numer", "typ", "wydzial", "proces", "nastepny")},
                    )
                )
            elif oczekujace.wyslano_at is None and oczekujace.blad == "Alert nieaktualny":
                oczekujace.blad = None
                oczekujace.nastepna_proba_at = teraz_utc()
        baza.flush()
        for wpis in baza.scalars(
            select(Powiadomienie)
            .where(
                Powiadomienie.wyslano_at.is_(None),
                Powiadomienie.nastepna_proba_at <= teraz_utc(),
            )
            .with_for_update(skip_locked=True)
            .limit(50)
        ):
            baza.scalar(select(Przyrzad).where(Przyrzad.id == wpis.przyrzad_id).with_for_update())
            swieze = pobierz_rejestr(baza, identyfikator=wpis.przyrzad_id)
            r = swieze[0] if swieze else None
            if (
                not r
                or not czy_wymaga_alertu(r, dzis_warszawa())
                or r["nastepny"] != wpis.dane["nastepny"]
            ):
                wpis.blad = "Alert nieaktualny"
                wpis.nastepna_proba_at = teraz_utc() + timedelta(days=365)
                continue
            wpis.dane = {k: r[k] for k in ("numer", "typ", "wydzial", "proces", "nastepny")}
            wiadomosc = EmailMessage()
            wiadomosc["From"] = os.environ["SMTP_FROM"]
            wiadomosc["To"] = "jakosc@aunde.pl"
            wiadomosc["Subject"] = "AUNDE: blokada przyrządu po terminie"
            wiadomosc["Message-ID"] = f"<{wpis.id}@metrologia.aunde.pl>"
            wiadomosc.set_content(
                "Przyrząd jest zablokowany z powodu przekroczenia terminu.\n\n"
                + "\n".join(f"{k}: {v}" for k, v in wpis.dane.items())
            )
            wpis.proby += 1
            try:
                with smtplib.SMTP(
                    os.environ["SMTP_HOST"],
                    int(os.getenv("SMTP_PORT", "587")),
                    timeout=15,
                ) as smtp:
                    smtp.starttls(context=ssl.create_default_context())
                    if os.getenv("SMTP_USER"):
                        smtp.login(os.environ["SMTP_USER"], os.environ["SMTP_PASSWORD"])
                    smtp.send_message(wiadomosc)
                wpis.wyslano_at, wpis.blad = teraz_utc(), None
            except (OSError, smtplib.SMTPException) as blad:
                wpis.blad = type(blad).__name__
                wpis.nastepna_proba_at = teraz_utc() + timedelta(
                    minutes=min(60, 2 ** min(wpis.proby, 6))
                )


# Uruchamia wyłącznie jawnie wybrane polecenie administracyjne; błędy kończą wykonanie bez pozornego sukcesu.
def uruchom() -> None:
    parser = argparse.ArgumentParser(description="Administracja AUNDE Metrologia")
    pod = parser.add_subparsers(dest="polecenie", required=True)
    konto = pod.add_parser("administrator")
    konto.add_argument("login")
    pod.add_parser("uprawnienia")
    pod.add_parser("powiadomienia")
    imp = pod.add_parser("import")
    imp.add_argument("plik")
    imp.add_argument("--zapisz", action="store_true")
    imp.add_argument("--raport", default="/dane/raport-importu.json")
    imp.add_argument("--konfiguracja", help="Lokalny JSON z listą polacz_numery, poza repozytorium")
    dane = parser.parse_args()
    if dane.polecenie == "administrator":
        utworz_administratora(dane.login)
    elif dane.polecenie == "uprawnienia":
        ustaw_uprawnienia()
    elif dane.polecenie == "powiadomienia":
        obsluz_powiadomienia()
    else:
        print(
            json.dumps(
                importuj_excel(
                    dane.plik,
                    dane.zapisz,
                    dane.raport,
                    set(
                        json.loads(Path(dane.konfiguracja).read_text(encoding="utf-8"))[
                            "polacz_numery"
                        ]
                    )
                    if dane.konfiguracja
                    else set(),
                ),
                ensure_ascii=False,
                indent=2,
            )
        )


if __name__ == "__main__":
    uruchom()
