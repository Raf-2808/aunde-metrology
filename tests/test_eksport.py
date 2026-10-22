from io import BytesIO
from zipfile import ZipFile
from openpyxl import load_workbook
from backend.eksport import zbuduj_excel


# Sprawdza oddzielenie obrazu QR od linku, dekodowanie, kolejność i ochronę tekstu przed formułami Excela.
def test_eksportu_z_qr():
    rekord = {
        "id": "test-id",
        "numer": "=1+1",
        "typ": "Miara",
        "rr": "Tak",
        "ostatni": "2026-01-01",
        "nastepny": "2027-01-01",
        "klasa": "B",
    }
    wynik = zbuduj_excel([rekord], {"test-id": "x" * 43}, "https://qr.example.test")
    with ZipFile(BytesIO(wynik)) as z:
        assert len([n for n in z.namelist() if n.startswith("xl/media/")]) == 1
        import zxingcpp
        from PIL import Image

        obraz = Image.open(
            BytesIO(z.read(next(n for n in z.namelist() if n.startswith("xl/media/"))))
        )
        for rozmiar in (obraz.size[0], 160):
            odczyt = zxingcpp.read_barcode(obraz.resize((rozmiar, rozmiar)))
            assert odczyt is not None
            assert odczyt.text == "https://qr.example.test/p/" + "x" * 43
    ksiazka = load_workbook(BytesIO(wynik))
    s = ksiazka.active
    assert s.cell(1, s.max_column).value == "Kod QR"
    assert [c.value for c in s[1]][:11] == [
        "Lp",
        "Numer",
        "Typ",
        "Stan",
        "Status",
        "Wydział",
        "Grupa",
        "R&R",
        "Ostatni",
        "Interwał",
        "Następny",
    ]
    assert s["B2"].data_type == "s" and s["B2"].value == "=1+1"
    assert s["S2"].hyperlink.target == "https://qr.example.test/p/" + "x" * 43
    assert s["S1"].value == "Adres karty QR"
    assert s["S2"].value == s["S2"].hyperlink.target
    assert s["T2"].value is None and s["T2"].hyperlink is None
    assert s._images[0].anchor._from.col == 19
    assert s._images[0].anchor._from.row == 1
    assert s["H2"].value == "Tak"
